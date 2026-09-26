"""Build the business-plan workbook (Excel) used as the 'plan' source of the DWH.

The workbook mimics how a bank's Planning department works:
  * GiaDinh        - bank-wide growth targets and seasonality (blue cells = inputs)
  * PhanBo_ChiNhanh - allocation of the bank target to branches (weights from 31/12/2023 actuals)
  * KeHoach_Thang  - monthly plan per branch & KPI (formulas) -> loaded by python/etl/load_plan.py

    python -m python.generator.build_plan_excel [--raw data/raw] [--out excel/templates/...]
"""
from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from python.common.config import ROOT, load_config
from python.common.layouts import LAYOUTS

FONT = "Arial"
BLUE, BLACK, GREEN = "0000FF", "000000", "008000"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
YEARS = [2024, 2025, 2026]
KPIS = [  # code, name, type
    ("HUY_DONG", "Tổng huy động cuối kỳ (VND)", "balance"),
    ("CASA", "Tiền gửi không kỳ hạn cuối kỳ (VND)", "balance"),
    ("DU_NO", "Dư nợ cho vay cuối kỳ (VND)", "balance"),
    ("KH_MOI", "Khách hàng mới trong tháng", "flow"),
    ("THE_TD", "Thẻ tín dụng phát hành trong tháng", "flow"),
]
GROWTH = {"HUY_DONG": [0.13, 0.14, 0.13], "CASA": [0.12, 0.14, 0.12], "DU_NO": [0.155, 0.17, 0.16],
          "KH_MOI": [0.10, 0.10, 0.08], "THE_TD": [0.12, 0.15, 0.10]}
# share of the annual increase reached by the end of each month (balances), and monthly share of flows
BAL_PROFILE = [0.04, 0.03, 0.14, 0.20, 0.27, 0.40, 0.45, 0.52, 0.63, 0.70, 0.80, 1.00]
FLOW_PROFILE = [0.075, 0.055, 0.085, 0.08, 0.085, 0.09, 0.08, 0.085, 0.09, 0.085, 0.09, 0.10]


def _read(raw: Path, entity: str) -> pd.DataFrame:
    f = raw / "history" / "202312" / f"DLB_{entity}_20231231.csv"
    df = pd.read_csv(f, sep="|", dtype=str, keep_default_na=False)
    return df.iloc[:-1] if df.iloc[-1, 0] == "TRL" else df


def base_values(raw: Path) -> pd.DataFrame:
    """Actual balances per branch at the migration date 31/12/2023."""
    acc = _read(raw, "CASA_ACCOUNT").drop_duplicates("account_no")
    cb = _read(raw, "CASA_BALANCE").merge(acc[["account_no", "branch_code"]], on="account_no")
    casa = cb.assign(v=cb.balance.astype(float)).groupby("branch_code").v.sum()
    td = _read(raw, "TERM_DEPOSIT").drop_duplicates("td_id")
    tb = _read(raw, "TD_BALANCE").merge(td[["td_id", "branch_code"]], on="td_id")
    tdv = tb.assign(v=tb.principal.astype(float)).groupby("branch_code").v.sum()
    ln = _read(raw, "LOAN").drop_duplicates("loan_id")
    lb = _read(raw, "LOAN_BALANCE").merge(ln[["loan_id", "branch_code"]], on="loan_id")
    lb = lb[lb.status.isin(["ACTIVE", "OVERDUE"])]
    loan = lb.assign(v=lb.outstanding_principal.astype(float)).groupby("branch_code").v.sum()
    cu = _read(raw, "CUSTOMER").drop_duplicates("cif")
    cust = cu.groupby("home_branch_code").size()
    ind = cu[cu.customer_type == "IND"].groupby("home_branch_code").size()
    br = pd.read_csv(ROOT / "data" / "reference" / "branches.csv", dtype=str)
    out = pd.DataFrame({"branch_code": br.branch_code, "branch_name": br.branch_name})
    out["CASA"] = out.branch_code.map(casa).fillna(0)
    out["HUY_DONG"] = out.CASA + out.branch_code.map(tdv).fillna(0)
    out["DU_NO"] = out.branch_code.map(loan).fillna(0)
    out["KH_MOI"] = out.branch_code.map(cust).fillna(0)
    out["THE_TD"] = out.branch_code.map(ind).fillna(0)
    return out


def _hdr(ws, row, headers, widths=None):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w


def build(raw: Path, out: Path) -> Path:
    base = base_values(raw)
    wb = Workbook()
    f_in = Font(name=FONT, color=BLUE)
    f_link = Font(name=FONT, color=GREEN)
    f_std = Font(name=FONT, color=BLACK)
    num = '#,##0;(#,##0);"-"'
    pct = '0.0%;(0.0%);"-"'

    # ------------------------------------------------------------------ HuongDan
    ws = wb.active
    ws.title = "HuongDan"
    lines = [
        ("KẾ HOẠCH KINH DOANH 2024–2026 — NGÂN HÀNG TMCP ĐÔNG LAM (hư cấu)", True),
        ("", False),
        ("Mục đích: giao chỉ tiêu theo chi nhánh/tháng. Sheet KeHoach_Thang được ETL nạp vào bảng dwh.fact_plan.", False),
        ("Quy trình: (1) Khối Kế hoạch nhập giả định ở sheet GiaDinh → (2) điều chỉnh hệ số ở PhanBo_ChiNhanh →", False),
        ("(3) KeHoach_Thang tự tính → (4) chạy: python -m python.etl.load_plan", False),
        ("", False),
        ("Quy ước màu:", True),
        ("  Chữ xanh dương, nền vàng nhạt = ô nhập liệu (được phép sửa)", False),
        ("  Chữ đen = công thức (không sửa)", False),
        ("  Chữ xanh lá = tham chiếu từ sheet khác", False),
        ("", False),
        ("Nguồn số liệu gốc: số dư thực tế ngày 31/12/2023 (file trích xuất core banking, thư mục data/raw/history/202312).", False),
        ("Chỉ tiêu số dư (HUY_DONG, CASA, DU_NO) là số dư cuối tháng; chỉ tiêu lưu lượng (KH_MOI, THE_TD) là số phát sinh trong tháng.", False),
        ("Kế hoạch năm sau lấy mục tiêu cuối năm trước làm điểm xuất phát (không lấy thực hiện), giống cách giao kế hoạch phổ biến.", False),
    ]
    for i, (t, b) in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=t).font = Font(name=FONT, bold=b, size=13 if i == 1 else 10)
    ws["A8"].font = Font(name=FONT, color=BLUE)
    ws["A8"].fill = INPUT_FILL
    ws["A10"].font = Font(name=FONT, color=GREEN)
    ws.column_dimensions["A"].width = 120

    # ------------------------------------------------------------------ GiaDinh
    g = wb.create_sheet("GiaDinh")
    g["A1"] = "GIẢ ĐỊNH TOÀN HÀNG"
    g["A1"].font = Font(name=FONT, bold=True, size=12)
    _hdr(g, 3, ["Mã KPI", "Chỉ tiêu", "Loại", "Thực hiện 31/12/2023", "Tăng trưởng 2024", "Tăng trưởng 2025",
                "Tăng trưởng 2026", "Mục tiêu 2024", "Mục tiêu 2025", "Mục tiêu 2026"],
         [12, 38, 10, 22, 14, 14, 14, 22, 22, 22])
    for i, (code, name, typ) in enumerate(KPIS):
        r = 4 + i
        g.cell(row=r, column=1, value=code).font = f_std
        g.cell(row=r, column=2, value=name).font = f_std
        g.cell(row=r, column=3, value="Số dư" if typ == "balance" else "Lưu lượng").font = f_std
        col = get_column_letter(4 + 0)
        c = g.cell(row=r, column=4, value=f"=SUM(PhanBo_ChiNhanh!{get_column_letter(3 + i)}5:{get_column_letter(3 + i)}{4 + len(base)})")
        c.font = f_link
        c.number_format = num
        if typ == "flow":
            c.comment = Comment("KPI lưu lượng: dùng quy mô khách hàng 31/12/2023 làm cơ sở; "
                                "mục tiêu năm = cơ sở × tỷ lệ ở cột D của bảng tỷ lệ bên dưới.", "Planning")
        for k, y in enumerate(YEARS):
            ci = g.cell(row=r, column=5 + k, value=GROWTH[code][k])
            ci.font, ci.fill, ci.number_format = f_in, INPUT_FILL, pct
        # targets
        if typ == "balance":
            g.cell(row=r, column=8, value=f"=D{r}*(1+E{r})")
            g.cell(row=r, column=9, value=f"=H{r}*(1+F{r})")
            g.cell(row=r, column=10, value=f"=I{r}*(1+G{r})")
        else:
            g.cell(row=r, column=8, value=f"=ROUND(D{r}*$D${16 + (i - 3)},0)")
            g.cell(row=r, column=9, value=f"=ROUND(H{r}*(1+F{r}),0)")
            g.cell(row=r, column=10, value=f"=ROUND(I{r}*(1+G{r}),0)")
        for k in range(3):
            cc = g.cell(row=r, column=8 + k)
            cc.font, cc.number_format = f_std, num
    g["A14"] = "Tỷ lệ cơ sở cho KPI lưu lượng năm 2024"
    g["A14"].font = Font(name=FONT, bold=True)
    _hdr(g, 15, ["Mã KPI", "Diễn giải", "", "Tỷ lệ/năm"])
    rates = [("KH_MOI", "Số KH mới năm 2024 = 16% số KH hiện hữu", 0.16),
             ("THE_TD", "Số thẻ TD phát hành năm 2024 = 7,5% số KH cá nhân hiện hữu", 0.075)]
    for i, (code, txt, v) in enumerate(rates):
        r = 16 + i
        g.cell(row=r, column=1, value=code).font = f_std
        g.cell(row=r, column=2, value=txt).font = f_std
        c = g.cell(row=r, column=4, value=v)
        c.font, c.fill, c.number_format = f_in, INPUT_FILL, pct
    g["A20"] = "Hồ sơ mùa vụ theo tháng"
    g["A20"].font = Font(name=FONT, bold=True)
    _hdr(g, 21, ["Tháng", "% tăng trưởng năm đạt được (số dư)", "", "% phân bổ trong năm (lưu lượng)"])
    for m in range(12):
        r = 22 + m
        g.cell(row=r, column=1, value=m + 1).font = f_std
        c1 = g.cell(row=r, column=2, value=BAL_PROFILE[m])
        c2 = g.cell(row=r, column=4, value=FLOW_PROFILE[m])
        for c in (c1, c2):
            c.font, c.fill, c.number_format = f_in, INPUT_FILL, pct
    g.cell(row=34, column=1, value="Tổng").font = Font(name=FONT, bold=True)
    g.cell(row=34, column=4, value="=SUM(D22:D33)").number_format = pct
    g.cell(row=35, column=1, value="Kiểm tra (phải = 100%)").font = Font(name=FONT, italic=True)
    g.cell(row=35, column=4, value='=IF(ABS(D34-1)<0.0001,"OK","SAI")')
    g["A37"] = ("Ghi chú: tăng trưởng mục tiêu tham chiếu định hướng tín dụng toàn ngành (2024: +15,08%; 2025: ~+18–19% "
                "theo NHNN) và tốc độ huy động thấp hơn tín dụng (2025: +14,11%). Nguồn: NHNN qua VnEconomy, "
                "Thị trường Tài chính Tiền tệ. Xem docs/02_calibration.md.")
    g["A37"].font = Font(name=FONT, italic=True, size=9)

    # ------------------------------------------------------------------ PhanBo_ChiNhanh
    p = wb.create_sheet("PhanBo_ChiNhanh")
    p["A1"] = "PHÂN BỔ CHỈ TIÊU THEO CHI NHÁNH"
    p["A1"].font = Font(name=FONT, bold=True, size=12)
    p["A2"] = "Cột C–G: thực hiện 31/12/2023 (nguồn: core banking). Cột H: hệ số điều chỉnh (đàm phán với chi nhánh)."
    p["A2"].font = Font(name=FONT, italic=True, size=9)
    heads = ["Mã CN", "Tên đơn vị"] + [f"TH 31/12/2023 {k[0]}" for k in KPIS] + ["Hệ số điều chỉnh"] + \
            [f"Tỷ trọng {k[0]}" for k in KPIS]
    _hdr(p, 4, heads, [9, 24] + [18] * 5 + [12] + [12] * 5)
    n = len(base)
    first, last = 5, 4 + n
    import numpy as np
    rng = np.random.default_rng(7)
    adj = np.round(rng.normal(1.0, 0.04, n), 2)
    for i, row in base.reset_index(drop=True).iterrows():
        r = first + i
        p.cell(row=r, column=1, value=row.branch_code).font = f_std
        p.cell(row=r, column=2, value=row.branch_name).font = f_std
        for k, (code, _, _) in enumerate(KPIS):
            c = p.cell(row=r, column=3 + k, value=float(row[code]))
            c.font, c.number_format = f_in, num
        c = p.cell(row=r, column=8, value=float(adj[i]))
        c.font, c.fill, c.number_format = f_in, INPUT_FILL, "0.00"
        for k in range(5):
            col = get_column_letter(3 + k)
            w = p.cell(row=r, column=9 + k,
                       value=f"={col}{r}*$H{r}/SUMPRODUCT({col}${first}:{col}${last},$H${first}:$H${last})")
            w.font, w.number_format = f_std, "0.00%"
    tr = last + 1
    p.cell(row=tr, column=2, value="Tổng").font = Font(name=FONT, bold=True)
    for k in range(10):
        if k == 5:
            continue
        col = get_column_letter(3 + k)
        c = p.cell(row=tr, column=3 + k, value=f"=SUM({col}{first}:{col}{last})")
        c.font = Font(name=FONT, bold=True)
        c.number_format = num if k < 5 else "0.00%"
    p.freeze_panes = "C5"

    # ------------------------------------------------------------------ KeHoach_Thang (long table)
    k = wb.create_sheet("KeHoach_Thang")
    _hdr(k, 1, ["plan_month", "branch_code", "kpi_code", "plan_value", "year", "month"], [13, 12, 11, 20, 8, 8])
    r = 2
    months = pd.period_range("2024-01", "2026-12", freq="M")
    for bi in range(n):
        br_row = first + bi
        for ki, (code, _, typ) in enumerate(KPIS):
            g_row = 4 + ki
            w_ref = f"PhanBo_ChiNhanh!${get_column_letter(9 + ki)}${br_row}"
            for per in months:
                yi = YEARS.index(per.year)
                tgt = f"GiaDinh!${get_column_letter(8 + yi)}${g_row}"
                start = f"GiaDinh!$D${g_row}" if yi == 0 else f"GiaDinh!${get_column_letter(7 + yi)}${g_row}"
                if typ == "balance":
                    prof = f"GiaDinh!$B${21 + per.month}"
                    f = f"=ROUND(({start}+({tgt}-{start})*{prof})*{w_ref},-6)"
                else:
                    prof = f"GiaDinh!$D${21 + per.month}"
                    f = f"=ROUND({tgt}*{prof}*{w_ref},0)"
                k.cell(row=r, column=1, value=per.end_time.date()).number_format = "yyyy-mm-dd"
                k.cell(row=r, column=2, value=base.branch_code.iat[bi])
                k.cell(row=r, column=3, value=code)
                c = k.cell(row=r, column=4, value=f)
                c.number_format = num
                k.cell(row=r, column=5, value=per.year)
                k.cell(row=r, column=6, value=per.month)
                r += 1
    tab = Table(displayName="tblKeHoach", ref=f"A1:F{r - 1}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    k.add_table(tab)
    k.freeze_panes = "A2"
    for row in k.iter_rows(min_row=2, max_row=r - 1):
        for c in row:
            c.font = f_std

    # ------------------------------------------------------------------ TongHop (pivot-like summary with SUMIFS)
    s = wb.create_sheet("TongHop", 1)
    s["A1"] = "TỔNG HỢP KẾ HOẠCH TOÀN HÀNG THEO THÁNG (tỷ VND / số lượng)"
    s["A1"].font = Font(name=FONT, bold=True, size=12)
    _hdr(s, 3, ["Tháng"] + [c for c, _, _ in KPIS], [13] + [16] * 5)
    for i, per in enumerate(months):
        rr = 4 + i
        s.cell(row=rr, column=1, value=per.end_time.date()).number_format = "mm/yyyy"
        for j, (code, _, typ) in enumerate(KPIS):
            div = "/1E9" if typ == "balance" else ""
            c = s.cell(row=rr, column=2 + j,
                       value=f'=SUMIFS(KeHoach_Thang!$D:$D,KeHoach_Thang!$A:$A,$A{rr},KeHoach_Thang!$C:$C,"{code}"){div}')
            c.number_format = "#,##0.0" if typ == "balance" else "#,##0"
            c.font = f_std
    s.freeze_panes = "B4"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--raw", default=None)
    ap.add_argument("--out", default=None)
    a = ap.parse_args()
    cfg = load_config()
    raw = Path(a.raw) if a.raw else cfg.paths.raw
    out = Path(a.out) if a.out else ROOT / "excel" / "templates" / "KE_HOACH_KINH_DOANH_2024_2026.xlsx"
    print(build(raw, out))


if __name__ == "__main__":
    main()
